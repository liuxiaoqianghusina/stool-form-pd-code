# Manuscript tables (Table 1-3) and supplementary tables (S1-S25) for the bowel-habits paper.
# Output: manuscript/Tables_main_and_supplementary.docx (three-line tables, Arial 9 pt) and
#         manuscript/Supplementary_Tables.xlsx (every supplementary table, full phenome-wide results included).
import os
import re

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

P = "/path/to/project/"
R = P + "results/"
OUT_DOCX = P + "manuscript/Tables_main_and_supplementary.docx"
OUT_XLSX = P + "manuscript/Supplementary_Tables.xlsx"

# ------------------------------------------------------------------ analytic cohort (same definition as scripts/prep.R)
co = pd.read_csv(P + "data/cohort.csv")
oc = pd.read_csv(P + "data/outcomes.csv")
keep = co.bmf_day.notna() & (co.excl_ibd == 0) & (co.excl_coeliac == 0) & (co.excl_crc == 0) & (co.excl_colectomy == 0) & (co.fu_end_days > 0)
co, oc = co[keep].reset_index(drop=True), oc[keep].reset_index(drop=True)
co["bmf_cat"] = np.select([co.bmf_day == 0, co.bmf_day == 1, co.bmf_day == 2], ["<1/day", "1/day", "2/day"], ">=3/day")
co["hard_cat"] = np.select([co.hard_stool == 0, co.hard_stool == 1, co.hard_stool >= 2], ["Never", "Sometimes", "Often+"], "")
co["loose_cat"] = np.select([co.loose_stool == 0, co.loose_stool == 1, co.loose_stool >= 2], ["Never", "Sometimes", "Often+"], "")
BMF = ["<1/day", "1/day", "2/day", ">=3/day"]
BMF_LAB = {"<1/day": "<1 per day", "1/day": "1 per day", "2/day": "2 per day", ">=3/day": "≥3 per day"}
FORM = ["Never", "Sometimes", "Often+"]
FORM_LAB = {"Never": "Never", "Sometimes": "Sometimes", "Often+": "Often or more"}

KEY = {"All-cause mortality": ["DEATH_all"], "CVD mortality": ["DEATH_CVD"],
       "Ischaemic heart disease": ["I20", "I21", "I22", "I23", "I24", "I25"], "Heart failure": ["I50"],
       "Atrial fibrillation": ["I48"], "Stroke": ["I60", "I61", "I63", "I64"], "Hypertension": ["I10"],
       "Type 2 diabetes": ["E11"], "Chronic kidney disease": ["N18"], "Depression": ["F32", "F33"],
       "Anxiety": ["F40", "F41"], "Dementia": ["F00", "F01", "F02", "F03", "G30"], "Parkinson's disease": ["G20"]}


def tt(name):
    """Event indicator, follow-up years and eligibility for a key outcome (mirrors outcome_tt in prep.R)."""
    codes = KEY[name]
    death = codes[0].startswith("DEATH_")
    m = oc[codes].to_numpy(dtype=float)
    prev = np.zeros(len(oc), bool) if death else (np.nan_to_num(m, nan=1e9) <= 0).any(axis=1)
    d = np.where(np.isnan(m).all(axis=1), np.nan, np.nanmin(np.where(np.isnan(m), np.inf, m), axis=1))
    fu = (co.fu_end_death_days if death else co.fu_end_days).to_numpy()
    event = (~np.isnan(d)) & (d > 0) & (d <= fu)
    time = np.where(event, d, fu) / 365.25
    return event, time, (~prev) & (time > 0)


def ev_py(name, by):
    e, t, s = tt(name)
    g = pd.DataFrame({"g": co[by].values, "e": e, "t": t})[s]
    a = g.groupby("g").agg(events=("e", "sum"), n=("e", "size"), py=("t", "sum"))
    a["ir"] = 1000 * a.events / a.py
    return a


def fmt_hr(hr, lo, hi):
    return f"{hr:.2f} ({lo:.2f}–{hi:.2f})"


def fmt_p(p):
    if pd.isna(p):
        return ""
    return "<0.001" if p < 0.001 else f"{p:.3f}"


# ------------------------------------------------------------------ Table 1
def n_pct(mask, base):
    k, n = int(mask[base].sum()), int(base.sum())
    return f"{k:,} ({100 * k / n:.1f})"


def mean_sd(x, base):
    v = x[base].dropna()
    return f"{v.mean():.1f} ({v.std():.1f})"


def med_iqr(x, base, dp=1):
    v = x[base].dropna()
    q1, md, q3 = v.quantile(.25), v.median(), v.quantile(.75)
    f = lambda z: f"{z:,.{dp}f}".replace("-", "−")
    sep = " to " if q1 < 0 or q3 < 0 else "–"  # avoid "−3.8–−0.1"
    return f"{f(md)} ({f(q1)}{sep}{f(q3)})"


groups = [("Overall", pd.Series(True, index=co.index))] + [(BMF_LAB[b], co.bmf_cat == b) for b in BMF]
T1_ROWS = [
    ("Participants, n", lambda b: f"{int(b.sum()):,}"),
    ("Age at questionnaire, years, mean (SD)", lambda b: mean_sd(co.age_q, b)),
    ("Female, n (%)", lambda b: n_pct(co.sex == 0, b)),
    ("White ethnicity, n (%)", lambda b: n_pct(co.ethnicity == "White", b)),
    ("Townsend deprivation index, median (IQR)", lambda b: med_iqr(co.tdi, b)),
    ("College or university degree, n (%)", lambda b: n_pct(co.degree == 1, b)),
    ("Smoking status, n (%)", None),
    ("  Never", lambda b: n_pct(co.smoking == 0, b)),
    ("  Previous", lambda b: n_pct(co.smoking == 1, b)),
    ("  Current", lambda b: n_pct(co.smoking == 2, b)),
    ("Alcohol daily or almost daily, n (%)", lambda b: n_pct(co.alcohol_freq == 1, b)),
    ("Physical activity, MET-min/week, median (IQR)", lambda b: med_iqr(co.met_min, b, 0)),
    ("Body mass index, kg/m², mean (SD)", lambda b: mean_sd(co.bmi, b)),
    ("Fresh fruit, pieces/day, median (IQR)", lambda b: med_iqr(co.fruit, b)),
    ("Vegetables, tablespoons/day, median (IQR)", lambda b: med_iqr(co.veg, b)),
    ("Wholemeal bread, n (%)", lambda b: n_pct(co.wholemeal_bread == 1, b)),
    ("Water, glasses/day, median (IQR)", lambda b: med_iqr(co.water_glasses, b)),
    ("Coffee, cups/day, median (IQR)", lambda b: med_iqr(co.coffee_cups, b)),
    ("Medication use, n (%)", None),
    ("  Laxatives", lambda b: n_pct(co.laxative_any == 1, b)),
    ("  Proton pump inhibitors", lambda b: n_pct(co.ppi == 1, b)),
    ("  Metformin", lambda b: n_pct(co.metformin == 1, b)),
    ("  Opioids", lambda b: n_pct(co.opioid == 1, b)),
    ("  Antidepressants", lambda b: n_pct(co.antidepressant == 1, b)),
    ("Conditions before the questionnaire, n (%)", None),
    ("  Diabetes", lambda b: n_pct(co.prev_diabetes == 1, b)),
    ("  Hypertension", lambda b: n_pct(co.prev_hypertension == 1, b)),
    ("  Ischaemic heart disease", lambda b: n_pct(co.prev_ihd == 1, b)),
    ("  Stroke", lambda b: n_pct(co.prev_stroke == 1, b)),
    ("  Depression", lambda b: n_pct(co.prev_depression == 1, b)),
    ("  Anxiety", lambda b: n_pct(co.prev_anxiety == 1, b)),
    ("  Irritable bowel syndrome", lambda b: n_pct(co.prev_ibs == 1, b)),
    ("  Chronic kidney disease", lambda b: n_pct(co.prev_ckd == 1, b)),
    ("  Cancer", lambda b: n_pct(co.prev_cancer == 1, b)),
    ("Stool form in the past 3 months, n (%)", None),
    ("  Hard/lumpy stools often or more", lambda b: n_pct(co.hard_stool >= 2, b)),
    ("  Loose/watery stools often or more", lambda b: n_pct(co.loose_stool >= 2, b)),
]
t1 = pd.DataFrame([[lab] + (["" for _ in groups] if f is None else [f(b) for _, b in groups]) for lab, f in T1_ROWS],
                  columns=["Characteristic"] + [g for g, _ in groups])

# ------------------------------------------------------------------ Table 2: stool form -> Parkinson's disease, dementia
k = pd.read_csv(R + "key_outcomes.csv")
dr = pd.read_csv(R + "dose_response.csv")
M1, M2 = "Model 1: age + sex", "Model 2: fully adjusted (main)"


def hr_cell(outcome, ex, lvl, analysis):
    d = k[(k.outcome == outcome) & (k.exposure == ex) & (k.level == lvl) & (k.analysis == analysis)]
    return "1.00 (reference)" if d.empty else fmt_hr(d.HR.iat[0], d.lo.iat[0], d.hi.iat[0])


rows = []
for outcome in ["Parkinson's disease", "Dementia"]:
    for ex, var, lab in [("hard_cat", "hard5", "Hard/lumpy stools"), ("loose_cat", "loose5", "Loose/watery stools")]:
        a = ev_py(outcome, ex)
        ptr = dr[(dr.exposure == var) & (dr.outcome == outcome)].p_trend.dropna()
        for i, lvl in enumerate(FORM):
            rows.append([outcome if i == 0 and ex == "hard_cat" else "", lab if i == 0 else "", FORM_LAB[lvl],
                         f"{int(a.at[lvl, 'events']):,}/{int(a.at[lvl, 'n']):,}", f"{a.at[lvl, 'ir']:.2f}",
                         hr_cell(outcome, ex, lvl, M1), hr_cell(outcome, ex, lvl, M2),
                         fmt_p(ptr.iat[0]) if (i == 0 and len(ptr)) else ""])
t2 = pd.DataFrame(rows, columns=["Outcome", "Exposure", "Frequency", "Events/N", "Rate per 1,000 PY", "HR (95% CI), model 1",
                                 "HR (95% CI), model 2", "P for trend†"])

# ------------------------------------------------------------------ Table 3: bowel movement frequency -> key outcomes
T3_OUT = ["Type 2 diabetes", "Hypertension", "Ischaemic heart disease", "Heart failure", "Atrial fibrillation", "Stroke",
          "Chronic kidney disease", "Depression", "Anxiety", "All-cause mortality", "CVD mortality"]
rows = []
for outcome in T3_OUT:
    a = ev_py(outcome, "bmf_cat")
    rows.append([outcome] + [f"{int(a.at[b, 'events']):,} ({a.at[b, 'ir']:.2f})" for b in BMF]
                + [hr_cell(outcome, "bmf_cat", b, M2) for b in ["<1/day", "2/day", ">=3/day"]])
t3 = pd.DataFrame(rows, columns=["Outcome"] + [f"Events (rate), {BMF_LAB[b]}" for b in BMF]
                  + ["HR (95% CI), <1 vs 1 per day", "HR (95% CI), 2 vs 1 per day", "HR (95% CI), ≥3 vs 1 per day"])

# ------------------------------------------------------------------ supplementary tables
S = {}
S["S1 Definitions"] = pd.DataFrame([
    ["Exposure", "Average bowel movements per day", "UKB field 21044 (digestive health questionnaire, 2017–2020)"],
    ["Exposure", "Least bowel movements per week", "UKB field 21043"],
    ["Exposure", "Hard/lumpy stools in the past 3 months", "UKB field 21033 (never, sometimes, often, most of the time, always)"],
    ["Exposure", "Loose/watery stools in the past 3 months", "UKB field 21034"],
    ["Outcome", "Incident diseases (phenome-wide)", "First occurrence of 3-character ICD-10 codes (Category 1712), chapters A–N"],
    ["Outcome", "Cancer", "Cancer registry, 3-character ICD-10 C codes (fields 40005/40006)"],
    ["Outcome", "Death", "Death register date (40000) and underlying cause (40001)"],
    ["Outcome", "Parkinson's disease", "ICD-10 G20 (first occurrence)"],
    ["Outcome", "Dementia", "ICD-10 F00–F03, G30 (first occurrence)"],
    ["Outcome", "Type 2 diabetes / hypertension / heart failure / atrial fibrillation", "E11 / I10 / I50 / I48"],
    ["Outcome", "Ischaemic heart disease / stroke / CKD", "I20–I25 / I60, I61, I63, I64 / N18"],
    ["Outcome", "Depression / anxiety", "F32–F33 / F40–F41"],
    ["Exclusion", "Before questionnaire", "IBD (K50, K51), coeliac/malabsorption (K90), colorectal cancer (C18–C20), colectomy or rectal excision (OPCS-4 H04–H11, H33)"],
    ["Follow-up", "Administrative censoring", "Hospital data: England 2023-03-31, Scotland 2022-09-30, Wales 2022-05-31; deaths 2024-09-30"],
    ["Covariate", "Model 1", "Age at questionnaire, sex"],
    ["Covariate", "Model 2", "Model 1 + ethnicity, Townsend index, degree, smoking, alcohol, MET-min/week, BMI, fruit, vegetables, wholemeal bread, fibre cereal, water, number of medications, laxatives, metformin, PPI, opioids, antidepressants, and prevalent diabetes, hypertension, IHD, stroke, depression, anxiety, IBS, CKD, cancer"],
    ["Covariate", "Model 3", "Model 2 + coffee and tea intake"],
], columns=["Type", "Item", "Definition / UKB field"])

pw = pd.read_csv(R + "phewas_model2.csv")
pw1 = pd.read_csv(R + "phewas_model1.csv")
LV = {("bmf_cat", "<1/day"): "<1 vs 1 per day", ("bmf_cat", "2/day"): "2 vs 1 per day", ("bmf_cat", ">=3/day"): "≥3 vs 1 per day",
      ("hard_cat", "Sometimes"): "Hard stools sometimes vs never", ("hard_cat", "Often+"): "Hard stools often+ vs never",
      ("loose_cat", "Sometimes"): "Loose stools sometimes vs never", ("loose_cat", "Often+"): "Loose stools often+ vs never",
      ("phenotype", "Constipation-type"): "Constipation-type vs normal", ("phenotype", "Diarrhoea-type"): "Diarrhoea-type vs normal",
      ("phenotype", "Mixed"): "Mixed vs normal"}


def tidy_pw(d):
    d = d.copy()
    d["Contrast"] = [LV.get((e, l), f"{e}: {l}") for e, l in zip(d.exposure, d.level)]
    d["HR (95% CI)"] = [fmt_hr(h, a, b) for h, a, b in zip(d.HR, d.lo, d.hi)]
    d["P"] = d.p.map(lambda v: f"{v:.2E}")
    d["FDR"] = d.fdr.map(lambda v: f"{v:.2E}")
    d["Outcome"] = [f"{c} {str(l).capitalize()}" if not str(c).startswith(("CR_", "DEATH_")) else str(l).capitalize()
                    for c, l in zip(d.code, d.label)]
    return d[["Contrast", "Outcome", "events_level", "events_total", "HR (95% CI)", "P", "FDR"]].rename(
        columns={"events_level": "Events in exposure group", "events_total": "Total events"})


sig = pw[pw.fdr < 0.05]
S["S2 PheWAS non-digestive"] = tidy_pw(sig[~sig.code.str.startswith("K")].sort_values(["exposure", "level", "p"]))
S["S3 PheWAS digestive"] = tidy_pw(sig[sig.code.str.startswith("K")].sort_values(["exposure", "level", "p"]))
S["S4a PheWAS all model 2"] = tidy_pw(pw.sort_values(["exposure", "level", "p"]))
S["S4b PheWAS all model 1"] = tidy_pw(pw1.sort_values(["exposure", "level", "p"]))

fg = pd.read_csv(R + "key_outcomes_finegray.csv")
ka = pd.concat([k, fg], ignore_index=True)
ka["Estimate (95% CI)"] = [fmt_hr(h, a, b) for h, a, b in zip(ka.HR, ka.lo, ka.hi)]
ka["Contrast"] = [LV.get((e, l), f"{e}: {l}") for e, l in zip(ka.exposure, ka.level)]
S["S5 Sensitivity analyses"] = ka[["outcome", "Contrast", "analysis", "events", "n", "Estimate (95% CI)", "p"]].rename(
    columns={"outcome": "Outcome", "analysis": "Analysis", "events": "Events", "n": "N", "p": "P"}).sort_values(["Outcome", "Contrast"])
ev = k[k.analysis == M2].dropna(subset=["evalue_point"])
S["S6 E-values"] = pd.DataFrame({"Outcome": ev.outcome, "Contrast": [LV.get((e, l), l) for e, l in zip(ev.exposure, ev.level)],
                                 "HR (95% CI)": [fmt_hr(h, a, b) for h, a, b in zip(ev.HR, ev.lo, ev.hi)],
                                 "E-value (point)": ev.evalue_point.round(2), "E-value (CI limit)": ev.evalue_ci.round(2)})
drs = dr.copy()
drs["HR (95% CI)"] = [("1.00 (reference)" if h == 1 else fmt_hr(h, a, b)) for h, a, b in zip(drs.HR, drs.lo, drs.hi)]
S["S7 Dose-response"] = drs[["exposure", "outcome", "level", "events", "n", "HR (95% CI)", "p", "p_trend",
                             "p_trend_below_1day", "p_trend_above_1day"]]

nh = pd.read_csv(R + "nhanes_validation.csv")
nh["Estimate (95% CI)"] = [fmt_hr(h, a, b) for h, a, b in zip(nh.estimate, nh.lo, nh.hi)]
S["S8 NHANES validation"] = nh[["outcome", "exposure", "level", "model", "Estimate (95% CI)", "p"]]
S["S9 NHANES treated PD"] = pd.read_csv(R + "nhanes_pd_meds.csv")
S["S10 MR instruments"] = pd.read_csv(R + "mr_instruments.csv")
mr = pd.read_csv(R + "mr_stoolfreq_finngen.csv")
mr["OR (95% CI)"] = [fmt_hr(h, a, b) for h, a, b in zip(mr.OR, mr.lo, mr.hi)]
S["S11 MR stool freq to outcomes"] = mr[["outcome", "instrument_set", "method", "nsnp", "OR (95% CI)", "pval", "Q_pval_ivw", "egger_intercept_p"]]
rv = pd.read_csv(R + "mr_reverse_pd_to_stoolfreq.csv")
rv["Beta (95% CI)"] = [f"{b:+.4f} ({b - 1.96 * s:+.4f} to {b + 1.96 * s:+.4f})" for b, s in zip(rv.b, rv.se)]
S["S12 Reverse MR PD to stool freq"] = rv[["exposure", "instrument_set", "method", "nsnp", "Beta (95% CI)", "pval", "Q_pval_ivw", "egger_intercept_p"]]
S["S13 Colocalisation"] = pd.read_csv(R + "coloc_stoolfreq_pd.csv").round(3)
q = pd.read_csv(R + "qsm_nigra.csv")
q["Beta, SD (95% CI)"] = [f"{b:+.3f} ({a:+.3f} to {c:+.3f})" for b, a, c in zip(q.beta_sd, q.lo, q.hi)]
S["S14 Brain iron QSM"] = q[["outcome", "exposure", "level", "model", "n", "Beta, SD (95% CI)", "p"]]
# S15: PD with stool form and frequency fitted jointly (common sample), plus Schoenfeld tests
ma = pd.read_csv(R + "pd_mutual_adjustment.csv")
ma["set"] = ma.analysis.map(lambda a: "single" if a.startswith("Single") else ("hl" if a.endswith("hard + loose") else "hlb"))
ma["var"] = ma.term.str.extract(r"^(hard_cat|loose_cat|bmf_cat)")[0]
ma["level"] = [t[len(v):] for t, v in zip(ma.term, ma["var"])]
ma["cell"] = [fmt_hr(h, a, b) for h, a, b in zip(ma.HR, ma.lo, ma.hi)]
VAR = {"hard_cat": "Hard/lumpy stools", "loose_cat": "Loose/watery stools", "bmf_cat": "Bowel movement frequency"}
LVL = {"Sometimes": "Sometimes", "Often+": "Often or more", "<1/day": "<1 per day", "2/day": "2 per day", ">=3/day": "≥3 per day"}
REF = {"hard_cat": "Never", "loose_cat": "Never", "bmf_cat": "1 per day"}
rows15 = []
for v in ["hard_cat", "loose_cat", "bmf_cat"]:
    rows15.append({"Exposure": VAR[v], "Category": REF[v] + " (reference)", "Single exposure": "1.00",
                   "Hard + loose stools": "1.00" if v != "bmf_cat" else "–", "Hard + loose stools + frequency": "1.00"})
    for lv in [l for l in LVL if l in set(ma.loc[ma["var"] == v, "level"])]:
        g = lambda s: ma.loc[(ma.set == s) & (ma["var"] == v) & (ma.level == lv), "cell"]
        rows15.append({"Exposure": "", "Category": LVL[lv], "Single exposure": g("single").iat[0],
                       "Hard + loose stools": g("hl").iat[0] if len(g("hl")) else "–",
                       "Hard + loose stools + frequency": g("hlb").iat[0]})
zp = pd.read_csv(R + "pd_schoenfeld.csv")
zp = zp[zp.term != "GLOBAL"].reset_index(drop=True)
for i, r in zp.iterrows():
    rows15.append({"Exposure": "Schoenfeld test" if i == 0 else "", "Category": VAR[r.exposure],
                   "Single exposure": f"χ² = {r.chisq:.2f}, P = {r.p:.2f}", "Hard + loose stools": "",
                   "Hard + loose stools + frequency": ""})
S["S15 PD mutual adjustment"] = pd.DataFrame(rows15)
S15_NOTE = (f"Hazard ratios (95% CI) for incident Parkinson's disease from model 2 in a common sample with complete data on "
            f"all three exposures (n = {int(ma.n.iat[0]):,}; {int(ma.events.iat[0])} events). Single exposure: each bowel "
            "habit modelled separately; the other columns include the listed exposures in the same model. Schoenfeld "
            "residual tests (exposure term) are from model 1, because the model 2 information matrix was singular "
            "owing to sparse covariate levels.")
# S16: frequency and non-PD outcomes (Table 3 moves to the supplement in the PD-focused version)
S["S16 Frequency other outcomes"] = t3
# S17: time origin at questionnaire (main) versus at recruitment (as in Zhang et al., Neuroepidemiology 2025)
to = pd.read_csv(R + "pd_time_origin_comparison.csv")
tc = pd.read_csv(R + "pd_time_origin_counts.csv").iloc[0]
DES = {"A": "Origin at questionnaire completion (main analysis)",
       "B": "Origin at recruitment; PD present at recruitment excluded",
       "B8": "Origin at recruitment; events in the first 8 years also excluded",
       "C": "Only PD diagnosed between recruitment and questionnaire"}
rows17 = []
for key, lab in DES.items():
    g = to[to.design.str.split(":").str[0] == key]
    cell = lambda ex, lv: fmt_hr(*g.loc[(g.exposure == ex) & (g.level == lv), ["HR", "lo", "hi"]].iloc[0])
    rows17.append({"Design": lab, "PD events": int(g.events.iloc[0]),
                   "Hard, sometimes": cell("hard_cat", "Sometimes"), "Hard, often or more": cell("hard_cat", "Often+"),
                   "Loose, sometimes": cell("loose_cat", "Sometimes"), "Loose, often or more": cell("loose_cat", "Often+")})
S["S17 Time origin comparison"] = pd.DataFrame(rows17)
S17_NOTE = (f"Hazard ratios (95% CI) for incident Parkinson's disease, model 2, never as reference. Bowel habits were "
            f"reported on the Digestive Health Questionnaire (2017–2020), a median of {tc.median_offset_y:.1f} years after "
            f"recruitment. With follow-up starting at recruitment, {int(tc.events_before_questionnaire)} of "
            f"{int(tc.events_recruit)} 'incident' cases ({100 * tc.events_before_questionnaire / tc.events_recruit:.0f}%) were "
            f"diagnosed before bowel habits were reported; {int(tc.events_before_q_after_8y)} such cases remained after "
            "excluding the first 8 years of follow-up.")
# S18: incremental predictive value of stool form over the conventional model (scripts/17_prediction.R)
pf = pd.read_csv(R + "pd_prediction_perf.csv").set_index("model")
pcv = pd.read_csv(R + "pd_prediction_cv.csv").set_index("model").C_cv
pen = pd.read_csv(R + "pd_prediction_enrich.csv")
pbt = pd.read_csv(R + "pd_prediction_boot.csv").set_index("stat")
psm = pd.read_csv(R + "pd_prediction_sample.csv").iloc[0]
MOD = {"base": "Conventional model", "freq": "+ bowel movement frequency", "form": "+ stool form (hard and loose stools)",
       "all": "+ stool form and frequency"}
rows18 = []
for k, lab in MOD.items():
    e10 = pen[(pen.model == k) & (pen.top_fraction == 0.10)]
    e05 = pen[(pen.model == k) & (pen.top_fraction == 0.05)]
    r = {"Model": lab, "C (apparent)": f"{(pf.C_base.iat[0] if k == 'base' else pf.loc[k, 'C_model']):.3f}",
         "C (10-fold CV)": f"{pcv[k]:.3f}" if k in pcv.index else "–",
         "ΔC vs conventional (95% CI)": "–" if k == "base" else f"{pf.loc[k, 'dC']:+.3f} ({pf.loc[k, 'dC_lo']:.3f} to {pf.loc[k, 'dC_hi']:.3f})",
         "LR χ² (df), P": "–" if k == "base" else f"{pf.loc[k, 'LR_chisq']:.1f} ({int(pf.loc[k, 'LR_df'])}), {pf.loc[k, 'LR_p']:.1e}",
         "Continuous NRI at 5 y": "–" if k == "base" else f"{pf.loc[k, 'nri']:.3f}",
         "Cases in top 10% risk, n (%)": "–" if e10.empty else f"{int(e10.cases_captured.iat[0])} ({e10.pct_cases_captured.iat[0]:.1f})",
         "Cases in top 5% risk, n (%)": "–" if e05.empty else f"{int(e05.cases_captured.iat[0])} ({e05.pct_cases_captured.iat[0]:.1f})"}
    rows18.append(r)
S["S18 Prediction"] = pd.DataFrame(rows18)
S18_NOTE = (f"n = {int(psm.n):,}; {int(psm.events)} incident Parkinson's disease cases. Conventional model: age (linear and "
            "quadratic), sex, smoking, coffee, physical activity, family history of Parkinson's disease, diabetes, depression, "
            "anxiety, hypotension (I95), urinary dysfunction (N31, N32, N39), sleep disorders (G47), diagnosed functional "
            "bowel disorder (K59) and laxative use. For the stool-form model, bootstrap 95% CIs (200 replicates) were "
            f"{pbt.loc['dC', 'lo']:.3f} to {pbt.loc['dC', 'hi']:.3f} for ΔC and {pbt.loc['nri', 'lo']:.3f} to "
            f"{pbt.loc['nri', 'hi']:.3f} for the NRI (events {pf.loc['form', 'nri_events']:.3f}, "
            f"{pbt.loc['nri_events', 'lo']:.3f} to {pbt.loc['nri_events', 'hi']:.3f}; non-events "
            f"{pf.loc['form', 'nri_nonevents']:.3f}, {pbt.loc['nri_nonevents', 'lo']:.3f} to {pbt.loc['nri_nonevents', 'hi']:.3f}). "
            "Five-year cumulative incidence in the top 5% of predicted risk: "
            f"{pen[(pen.model == 'base') & (pen.top_fraction == 0.05)].risk5_top_km.iat[0]:.2f}% (conventional) and "
            f"{pen[(pen.model == 'form') & (pen.top_fraction == 0.05)].risk5_top_km.iat[0]:.2f}% (with stool form). "
            "CV, cross-validated; LR, likelihood ratio; NRI, net reclassification improvement.")

# ------------------------------------------------------------------ S19-S25: tier-1/tier-2 analyses (18_tier1.R, 19_tier2.R)
def hr_fmt(df):
    return [("1.00 (reference)" if pd.isna(p) else fmt_hr(h, a, b)) for h, a, b, p in zip(df.HR, df.lo, df.hi, df.p)]


TERM = {"hard_catSometimes": "Hard stools, sometimes", "hard_catOften+": "Hard stools, often or more",
        "loose_catSometimes": "Loose stools, sometimes", "loose_catOften+": "Loose stools, often or more"}
nc = pd.read_csv(R + "tier1_negative_control.csv"); ncc = pd.read_csv(R + "tier1_negative_control_counts.csv")
NCL = {"abd_pain_freqLess than weekly": ("Abdominal pain frequency", "Less than weekly"),
       "abd_pain_freqWeekly or more": ("Abdominal pain frequency", "Weekly or more"),
       "abd_pain_nowYes": ("Current abdominal pain", "Yes"), "distension_nowYes": ("Current abdominal distension", "Yes")}
s19 = []
for t, (e, lv) in NCL.items():
    a = nc[(nc.term == t) & (nc.analysis == "Single exposure")].iloc[0]; b = nc[(nc.term == t) & (nc.analysis != "Single exposure")].iloc[0]
    exn = {"Abdominal pain frequency": "abd_pain_freq", "Current abdominal pain": "abd_pain_now", "Current abdominal distension": "distension_now"}[e]
    c = ncc[(ncc.exposure == exn) & (ncc.level == lv)].iloc[0]
    s19.append({"Exposure": e, "Level (vs never/no)": lv, "Participants": f"{int(c.n):,}", "PD events": int(c.events),
                "HR (95% CI)": fmt_hr(a.HR, a.lo, a.hi), "P": fmt_p(a.p),
                "HR additionally adjusted for stool form": fmt_hr(b.HR, b.lo, b.hi), "P ": fmt_p(b.p)})
S["S19 Negative control exposures"] = pd.DataFrame(s19)

# S20: further sensitivity analyses (outcome definitions, medication, missing data)
de = pd.read_csv(R + "tier1_pd_definitions.csv"); me = pd.read_csv(R + "tier2_medication.csv")
md = pd.read_csv(R + "tier2_missing_data.csv") if os.path.exists(R + "tier2_missing_data.csv") else pd.DataFrame(columns=de.columns)
slc = pd.read_csv(R + "sens_loose_causes.csv")
s20 = pd.concat([de, me, slc, md], ignore_index=True)
s20 = s20[s20.term.isin(TERM)].copy()
s20["Exposure"] = s20.term.map(TERM); s20["HR (95% CI)"] = hr_fmt(s20)
s20 = s20.pivot_table(index=["analysis", "events"], columns="Exposure", values="HR (95% CI)", aggfunc="first").reset_index()
s20 = s20.rename(columns={"analysis": "Analysis", "events": "PD events"})[["Analysis", "PD events"] + list(TERM.values())]
S["S20 Further sensitivity"] = s20

# S21: transit spectrum score
ts = pd.read_csv(R + "tier1_transit_score.csv"); tsc = pd.read_csv(R + "tier1_transit_score_counts.csv")
tsi = pd.read_csv(R + "tier1_transit_info.csv").iloc[0]
SC = {"score_cat<=-3": ([-4, -3], "≤−3 (mostly loose)"), "score_cat-2": ([-2], "−2"), "score_cat-1": ([-1], "−1"),
      "score_cat0": ([0], "0 (reference)"), "score_cat1": ([1], "1"), "score_cat2": ([2], "2"), "score_cat>=3": ([3, 4], "≥3 (mostly hard)")}
s21 = []
for t, (vals, lab) in SC.items():
    r = ts[(ts.analysis == "categorical") & (ts.term == t)].iloc[0]; c = tsc[tsc.score.isin(vals)]
    s21.append({"Transit spectrum score": lab, "Participants": f"{int(c.n.sum()):,}", "PD events": int(c.events.sum()),
                "HR (95% CI)": "1.00 (reference)" if t == "score_cat0" else fmt_hr(r.HR, r.lo, r.hi), "P": fmt_p(r.p)})
for an in ["per 1-point increase", "per 1-point increase, mixed pattern excluded"]:
    r = ts[ts.analysis == an].iloc[0]
    s21.append({"Transit spectrum score": "Per 1-point increase" + (" (mixed pattern excluded)" if "mixed" in an else ""),
                "Participants": f"{int(r.n):,}", "PD events": int(r.events), "HR (95% CI)": fmt_hr(r.HR, r.lo, r.hi), "P": fmt_p(r.p)})
S["S21 Transit spectrum score"] = pd.DataFrame(s21)
S21_NOTE = (f"Score = hard-stool frequency level minus loose-stool frequency level (each 0 = never to 4 = always). Model 2. "
            f"P for non-linearity (natural spline, 3 df, versus linear) = {tsi.p_nonlinearity:.2f}; "
            f"{int(tsi.n_mixed):,} participants reported both hard and loose stools often or more (mixed pattern).")

# S22: subgroups
sg = pd.read_csv(R + "tier1_subgroups.csv"); sg = sg[sg.term.str.contains("Often")]
SGL = {"sg_sex": "Sex", "sg_age": "Age", "sg_smoke": "Smoking", "sg_coffee": "Coffee", "sg_bmi": "Body mass index",
       "sg_lax": "Laxative use", "sg_dep": "Depression history"}
s22 = []
for g, gl in SGL.items():
    h = sg[(sg.subgroup == g) & (sg.exposure == "hard_cat")]; l = sg[(sg.subgroup == g) & (sg.exposure == "loose_cat")]
    for lv in h.lvl:
        a = h[h.lvl == lv].iloc[0]; b = l[l.lvl == lv].iloc[0]
        s22.append({"Subgroup": gl, "Level": lv.replace(">=", "≥").replace("kg/m2", "kg/m²"), "Participants": f"{int(a.n):,}",
                    "PD events": int(a.events), "Hard stools often+ HR (95% CI)": fmt_hr(a.HR, a.lo, a.hi),
                    "P-interaction (hard)": f"{a.p_interaction:.2g}", "Loose stools often+ HR (95% CI)": fmt_hr(b.HR, b.lo, b.hi),
                    "P-interaction (loose)": f"{b.p_interaction:.2g}"})
S["S22 Subgroups"] = pd.DataFrame(s22)

# S23: absolute risk
aj = pd.read_csv(R + "tier1_cuminc_aj.csv"); ar = pd.read_csv(R + "tier1_absolute_risk.csv")
LVN = {"Never": "Never", "Sometimes": "Sometimes", "Often+": "Often or more"}
s23 = []
for pop, pre in [("All participants", ""), ("Aged ≥65 years", "age65 ")]:
    for e, en in [("hard", "Hard stools"), ("loose", "Loose stools")]:
        for lv in ["Never", "Sometimes", "Often+"]:
            a = ar[ar.group == f"{pre}{e} {lv}"].iloc[0]
            c = aj[(aj.exposure == f"{e}_cat") & (aj.level == lv)].iloc[0] if pre == "" else None
            s23.append({"Population": pop, "Exposure": en, "Level": LVN[lv],
                        "Crude 5-y cumulative incidence, % (95% CI)": "–" if c is None else f"{c.cuminc5:.2f} ({c.lo:.2f}–{c.hi:.2f})",
                        "Standardized 5-y risk, % (95% CI)": f"{a.risk5_pct:.2f} ({a.risk5_lo:.2f}–{a.risk5_hi:.2f})",
                        "Risk difference per 1,000 (95% CI)": "0 (reference)" if lv == "Never" else f"{a.rd_per1000:+.1f} ({a.rd_lo:+.1f} to {a.rd_hi:+.1f})"})
S["S23 Absolute risk"] = pd.DataFrame(s23)

# S24: stool form jointly with other prodromal markers
jt = pd.read_csv(R + "tier2_joint.csv"); jc = pd.read_csv(R + "tier2_joint_counts.csv"); ji = pd.read_csv(R + "tier2_joint_cuminc.csv")
wi = pd.read_csv(R + "tier2_within_prodromal.csv"); it = pd.read_csv(R + "tier2_interaction.csv").iloc[0]
s24 = []
for form in ["Loose often+", "Neither", "Hard often+"]:
    for pm in ["No other prodromal marker", "Any other prodromal marker"]:
        lab = f"{form} | {pm}"; r = jt[jt.term == "joint" + lab].iloc[0]; c = jc[jc.joint == lab].iloc[0]; i = ji[ji.joint == lab].iloc[0]
        s24.append({"Stool form": {"Loose often+": "Loose stools often or more", "Neither": "Neither",
                                   "Hard often+": "Hard stools often or more"}[form],
                    "Other prodromal markers": "None" if pm.startswith("No") else "≥1", "Participants": f"{int(c.n):,}",
                    "PD events": int(c.events), "HR (95% CI)": "1.00 (reference)" if pd.isna(r.p) else fmt_hr(r.HR, r.lo, r.hi),
                    "5-y cumulative incidence, % (95% CI)": f"{i.cuminc5:.2f} ({i.lo:.2f}–{i.hi:.2f})"})
S["S24 Joint prodromal markers"] = pd.DataFrame(s24)
S24_NOTE = ("Other prodromal markers recorded before the questionnaire: hypotension (I95), urinary dysfunction (N31, N32, N39), "
            "sleep disorders (G47), depression or anxiety (F32, F33, F40, F41). Participants reporting both hard and loose stools "
            "often or more were excluded. Within strata, hard stools often+ HR "
            f"{fmt_hr(*wi[(wi.analysis.str.startswith('No')) & (wi.term == 'hard_catOften+')][['HR', 'lo', 'hi']].iloc[0])} (no marker) and "
            f"{fmt_hr(*wi[(wi.analysis.str.startswith('Any')) & (wi.term == 'hard_catOften+')][['HR', 'lo', 'hi']].iloc[0])} (≥1 marker); "
            f"loose stools often+ {fmt_hr(*wi[(wi.analysis.str.startswith('No')) & (wi.term == 'loose_catOften+')][['HR', 'lo', 'hi']].iloc[0])} and "
            f"{fmt_hr(*wi[(wi.analysis.str.startswith('Any')) & (wi.term == 'loose_catOften+')][['HR', 'lo', 'hi']].iloc[0])}. "
            f"P for multiplicative interaction {it.p_mult_hard:.2f} (hard) and {it.p_mult_loose:.2f} (loose); relative excess risk "
            f"due to interaction for hard stools x any marker {it.RERI:.2f} ({it.RERI_lo:.2f} to {it.RERI_hi:.2f}).")

# S25: accelerometry sub-study
ah = pd.read_csv(R + "tier2_accel_hr.csv"); ad = pd.read_csv(R + "tier2_accel_discrimination.csv")
ai = pd.read_csv(R + "tier2_accel_info.csv").iloc[0]; aj2 = pd.read_csv(R + "tier2_accel_joint.csv"); ajc = pd.read_csv(R + "tier2_accel_joint_counts.csv")
ACT = {"acc_sd": "Average acceleration, per SD", **TERM}
s25 = [{"Analysis": an, "Term": ACT.get(r.term, r.term), "HR (95% CI)": fmt_hr(r.HR, r.lo, r.hi), "P": fmt_p(r.p)}
       for an, d in ah.groupby("analysis", sort=False) for _, r in d.iterrows()]
JL = {"interaction(acc_low, hard_stool >= 2)1.FALSE": "Lowest acceleration quintile, hard stools not often",
      "interaction(acc_low, hard_stool >= 2)0.TRUE": "Higher acceleration, hard stools often or more",
      "interaction(acc_low, hard_stool >= 2)1.TRUE": "Lowest acceleration quintile and hard stools often or more"}
s25 += [{"Analysis": "Joint categories (reference: higher acceleration, hard stools not often)", "Term": JL[r.term],
         "HR (95% CI)": fmt_hr(r.HR, r.lo, r.hi), "P": fmt_p(r.p)} for _, r in aj2.iterrows()]
s25 += [{"Analysis": "Discrimination (C index)", "Term": r.model, "HR (95% CI)": f"C = {r.C:.3f}" + ("" if pd.isna(r.dC_vs_previous) else
         f"; ΔC {r.dC_vs_previous:+.3f} ({r.lo:.3f} to {r.hi:.3f}) {r.comparison}"), "P": ""} for _, r in ad.iterrows()]
S["S25 Accelerometry"] = pd.DataFrame(s25)
S25_NOTE = (f"Participants with valid wrist accelerometry (good wear time and calibration; 2013–2015): n = {int(ai.n):,}; "
            f"{int(ai.events)} incident PD cases; median {ai.median_gap_y:.1f} years between accelerometry and the questionnaire. "
            "Hazard ratios from model 2; discrimination models include age, age squared and sex.")

# ------------------------------------------------------------------ abbreviations for every table footnote
CI, HR, PDA = "CI, confidence interval", "HR, hazard ratio", "PD, Parkinson's disease"
ICD = "ICD-10, International Classification of Diseases, 10th revision"
MRA = ("IVW, inverse-variance weighted; MAPT, microtubule-associated protein tau; MR, Mendelian randomization; MR-PRESSO, "
       "Mendelian Randomization Pleiotropy RESidual Sum and Outlier; nsnp, number of instruments; Q_pval_ivw, P value of "
       "Cochran's Q for the IVW estimate; egger_intercept_p, P value of the MR-Egger intercept")
ABBR = {
    "T1": "IBD, inflammatory bowel disease; IQR, interquartile range; MET, metabolic equivalent of task; SD, standard deviation.",
    "T2": f"BMI, body mass index; {CI}; CKD, chronic kidney disease; {HR}; IBS, irritable bowel syndrome; PPI, proton pump inhibitor; PY, person-years.",
    "T3": f"{CI}; CVD, cardiovascular disease; {HR}.",
    "S1": ("BMI, body mass index; CKD, chronic kidney disease; IBD, inflammatory bowel disease; IBS, irritable bowel syndrome; "
           f"{ICD}; IHD, ischaemic heart disease; MET, metabolic equivalent of task; OPCS-4, Office of Population Censuses and "
           "Surveys Classification of Interventions and Procedures, version 4; PPI, proton pump inhibitor; UKB, UK Biobank."),
    "S2": f"Outcomes are three-character {ICD.split(', ', 1)[0]} codes, cancer registry sites or causes of death. {CI}; FDR, false discovery rate (Benjamini–Hochberg); {HR}; {ICD}.",
    "S5": f"{CI}; CVD, cardiovascular disease; IBS, irritable bowel syndrome. Estimates are hazard ratios.",
    "S6": (f"{CI}; CVD, cardiovascular disease; {HR}. The E-value is the minimum strength of association, on the risk-ratio "
           "scale, that an unmeasured confounder would need with both exposure and outcome to explain away the estimate (point) "
           "or its CI limit."),
    "S7": (f"{CI}; {HR}. p_trend, P for linear trend across levels; p_trend_below_1day and p_trend_above_1day, P for trend "
           "across frequency categories below and above one bowel movement per day; HR_per_level, hazard ratio per one-level increase."),
    "S8": f"BSFS, Bristol Stool Form Scale; {CI}; OR, odds ratio. Estimates are hazard ratios for mortality and odds ratios for cross-sectional outcomes.",
    "S9": ("BSFS, Bristol Stool Form Scale; COMT, catechol-O-methyltransferase; MAO-B, monoamine oxidase B; OR, odds ratio; "
           f"{PDA}; lo and hi, lower and upper limits of the 95% confidence interval; cases_total, number of drug users."),
    "S10": ("chr, chromosome; pos, position (GRCh38); beta, effect on stool frequency per effect allele; standard_error, standard "
            "error of beta; F, F statistic (beta²/standard error²)."),
    "S11": f"BDNF, brain-derived neurotrophic factor; {CI}; {MRA}; OR, odds ratio.",
    "S12": f"{CI}; {MRA}; {PDA}; R12, FinnGen release 12. Beta is the change in stool frequency per unit log-odds of genetic liability to PD.",
    "S13": ("chr, chromosome; pos, position (GRCh38); n_snps, number of variants tested; PP_H0 to PP_H4, posterior probabilities of "
            "no association (H0), association with stool frequency only (H1) or PD only (H2), association with both through distinct "
            "causal variants (H3) or through a shared causal variant (H4); top_shared_snp, variant with the highest posterior "
            f"probability of being shared; min_p_pd, smallest P value for PD in the region; {PDA}."),
    "S14": f"{CI}; QSM, quantitative susceptibility mapping; SD, standard deviation; t2s, T2* (effective transverse relaxation time).",
    "S15": f"{CI}; {PDA}.",
    "S16": f"{CI}; CVD, cardiovascular disease; {HR}.",
    "S17": f"{CI}; {PDA}.",
    "S18": f"{CI}; {PDA}; ΔC, difference in C index.",
    "S19": f"{CI}; {HR}; {PDA}.",
    "S20": (f"{CI}; E05, thyrotoxicosis; G21, secondary parkinsonism; G23, other degenerative diseases of basal ganglia; K52, other "
            f"non-infective gastroenteritis and colitis; K86, other diseases of pancreas (ICD-10 codes); {PDA}; PPI, proton pump inhibitor."),
    "S21": f"{CI}; {HR}; {PDA}.",
    "S22": f"{CI}; {HR}; {PDA}.",
    "S23": f"{CI}.",
    "S24": f"{CI}; {HR}; {PDA}.",
    "S25": f"{CI}; {HR}; {PDA}; SD, standard deviation.",
}
ABBR["S3"] = ABBR["S4"] = ABBR["S2"]


def table_key(label):
    """'Table 1', 'Table 1. ...', 'S4a PheWAS ...' or 'Supplementary Table S15. ...' -> 'T1' / 'S4' / 'S15'."""
    m = re.search(r"(?:Supplementary Table )?(S\d+)|Table (\d+)", label)
    return m.group(1) if m.group(1) else "T" + m.group(2)


# ------------------------------------------------------------------ Excel
with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as xw:
    for name, d in [("Table 1", t1), ("Table 2", t2), ("Table 3", t3)] + list(S.items()):
        d.to_excel(xw, sheet_name=name[:31], index=False)
        ws = xw.sheets[name[:31]]
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, max(len(str(c.value or "")) for c in col[:200]) + 2))
        ws.cell(row=len(d) + 3, column=1, value="Abbreviations: " + ABBR[table_key(name)])

# ------------------------------------------------------------------ Word (three-line tables, Arial 9 pt)
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(9)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")


def border(cell, edges):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for e in ("top", "bottom", "left", "right"):
        el = OxmlElement(f"w:{e}")
        if e in edges:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "8"); el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tcPr.append(b)


def add_table(title, d, note, landscape=False, widths=None, abbr_key=None):
    sec = doc.sections[-1] if len(doc.paragraphs) == 0 else doc.add_section()
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        if sec.page_width < sec.page_height:
            sec.page_width, sec.page_height = sec.page_height, sec.page_width
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
        if sec.page_width > sec.page_height:
            sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Cm(1.6))
    p = doc.add_paragraph(); r = p.add_run(title); r.bold = True; r.font.size = Pt(10)
    tbl = doc.add_table(rows=len(d) + 1, cols=d.shape[1])
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    usable = (sec.page_width - sec.left_margin - sec.right_margin) / 360000  # EMU -> cm
    if not widths:  # first column wider, the rest share the remaining width
        first = min(7.0, usable * 0.28)
        widths = [first] + [(usable - first) / (d.shape[1] - 1)] * (d.shape[1] - 1)
    for j, h in enumerate(d.columns):
        c = tbl.cell(0, j); c.text = ""; run = c.paragraphs[0].add_run(str(h)); run.bold = True; run.font.size = Pt(8.5)
        border(c, ("top", "bottom"))
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
    for i, row in enumerate(d.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            c = tbl.cell(i, j); c.text = ""
            run = c.paragraphs[0].add_run("" if pd.isna(v) else str(v)); run.font.size = Pt(8.5)
            if j > 0:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            c.paragraphs[0].paragraph_format.space_before = Pt(1)
            border(c, ("bottom",) if i == len(d) else ())
    for j, w in enumerate(widths):
        tbl.columns[j].width = Cm(w)
        for row in tbl.rows:
            row.cells[j].width = Cm(w)
    note = (note + " " if note else "") + "Abbreviations: " + ABBR[abbr_key or table_key(title)]
    n = doc.add_paragraph(); rn = n.add_run(note); rn.italic = True; rn.font.size = Pt(8)


add_table("Table 1. Characteristics of participants by average bowel movement frequency", t1,
          "Values are n (%), mean (SD) or median (IQR). Lifestyle and medication data are from the UK Biobank baseline "
          "assessment (2006–2010); conditions are those recorded before the digestive health questionnaire (2017–2020). "
          "Participants with IBD, coeliac disease or malabsorption, colorectal cancer or prior colectomy were excluded.",
          widths=[6.2, 2.5, 2.4, 2.5, 2.4, 2.4])
add_table("Table 2. Stool form and incident Parkinson's disease and dementia", t2,
          "Model 1: age and sex. Model 2: additionally adjusted for ethnicity, Townsend "
          "index, education, smoking, alcohol, physical activity, BMI, diet (fruit, vegetables, wholemeal bread, fibre cereal), "
          "water intake, number of medications, laxatives, metformin, PPIs, opioids, antidepressants and prevalent diabetes, "
          "hypertension, ischaemic heart disease, stroke, depression, anxiety, IBS, CKD and cancer. †Linear trend across the "
          "five response levels (never, sometimes, often, most of the time, always), model 2.", landscape=True)
add_table("Table 3. Bowel movement frequency and incident cardiometabolic disease, mental disorders and mortality", t3,
          "Events (rate per 1,000 person-years) by category; hazard ratios from model 2 with 1 bowel movement per day as the "
          "reference. Covariates as in Table 2.", landscape=True)
for name in ["S1 Definitions", "S2 PheWAS non-digestive", "S6 E-values", "S8 NHANES validation", "S9 NHANES treated PD",
             "S11 MR stool freq to outcomes", "S12 Reverse MR PD to stool freq", "S13 Colocalisation"]:
    d = S[name]
    if name == "S11 MR stool freq to outcomes":
        d = d[d.method.isin(["Inverse variance weighted", "MR Egger", "Weighted median", "MR-PRESSO Outlier-corrected"])]
    add_table(f"Supplementary Table {name.split()[0]}. {' '.join(name.split()[1:])}", d.round(4),
              "Full versions of all supplementary tables, including the complete phenome-wide results (S4) and every "
              "sensitivity analysis (S5), are provided in Supplementary_Tables.xlsx.", landscape=True)
add_table("Supplementary Table S15. Stool form and bowel movement frequency fitted jointly for incident Parkinson's disease",
          S["S15 PD mutual adjustment"], S15_NOTE, widths=[3.6, 3.4, 3.6, 3.6, 3.8])
add_table("Supplementary Table S17. Choice of time origin and the association of stool form with Parkinson's disease",
          S["S17 Time origin comparison"], S17_NOTE, landscape=True, widths=[8.0, 2.0, 3.4, 3.4, 3.4, 3.4])
add_table("Supplementary Table S18. Incremental predictive value of stool form for incident Parkinson's disease",
          S["S18 Prediction"], S18_NOTE, landscape=True, widths=[5.2, 2.0, 2.0, 3.6, 3.2, 2.2, 2.6, 2.6])
add_table("Supplementary Table S19. Negative-control exposures from the Digestive Health Questionnaire and incident Parkinson's disease",
          S["S19 Negative control exposures"], "Model 2. Reference: never (frequency) or no (current symptoms).", landscape=True)
add_table("Supplementary Table S20. Further sensitivity analyses for stool form and incident Parkinson's disease",
          S["S20 Further sensitivity"], "Hazard ratios (95% CI), model 2 (multiple imputation and complete-case models include the "
          "same covariates with missing values imputed or excluded instead of median-filled). Reference: never.", landscape=True,
          widths=[7.2, 1.6, 3.6, 3.6, 3.6, 3.6])
add_table("Supplementary Table S21. Transit spectrum score and incident Parkinson's disease", S["S21 Transit spectrum score"], S21_NOTE)
add_table("Supplementary Table S22. Stool form and incident Parkinson's disease by subgroup", S["S22 Subgroups"],
          "Model 2; P for interaction from likelihood-ratio tests of exposure-by-subgroup product terms.", landscape=True)
add_table("Supplementary Table S23. Absolute 5-year risk of Parkinson's disease by stool form", S["S23 Absolute risk"],
          "Crude: Aalen–Johansen cumulative incidence with death as a competing event. Standardized: mean predicted 5-year risk "
          "from model 2 with every participant assigned each category; 95% CIs from 200 bootstrap replicates.", landscape=True)
add_table("Supplementary Table S24. Stool form jointly with other prodromal markers and incident Parkinson's disease",
          S["S24 Joint prodromal markers"], S24_NOTE, landscape=True)
add_table("Supplementary Table S25. Accelerometry sub-study: stool form, average acceleration and incident Parkinson's disease",
          S["S25 Accelerometry"], S25_NOTE, landscape=True, widths=[6.0, 6.5, 8.0, 2.0])
doc.save(OUT_DOCX)
print("saved:", OUT_DOCX, "|", OUT_XLSX)
print(t1.head(8).to_string(index=False)); print(t2.to_string(index=False)); print(t3.to_string(index=False))


# ================================================================== submission package
# Supplementary Information (Word: Supplementary Figs 1-3 with legends + Supplementary Tables 1-16) and Supplementary Data
# (Excel: large tables, Supplementary Data 1-10). Numbering follows first citation in the English manuscript and is read
# from manuscript/supplement_map.json (written by 24_renumber_supplement_text.py).
import json

from docx.shared import Cm as _Cm
from openpyxl import Workbook
from openpyxl.styles import Font

MAP = json.load(open(P + "manuscript/supplement_map.json"))
SKEY = {k.split()[0][1:]: k for k in S}                      # "15" -> "S15 PD mutual adjustment", "4a" -> "S4a ..."
PDn = "Parkinson's disease"
TITLES = {
    "1": "Definitions of exposures, outcomes and covariates",
    "2": "Phenome-wide associations at FDR < 0.05: outcomes outside the digestive system",
    "3": "Phenome-wide associations at FDR < 0.05: digestive outcomes",
    "4a": "Complete phenome-wide results, fully adjusted model (model 2)",
    "4b": "Complete phenome-wide results, age- and sex-adjusted model (model 1)",
    "5": "Model ladder and sensitivity analyses for 14 key outcomes",
    "6": "E-values for key associations",
    "7": "Dose–response analyses for stool form and bowel movement frequency",
    "8": "Bowel habits and mortality and cross-sectional outcomes in NHANES 2005–2010",
    "9": "Stool form and antiparkinsonian drug use in NHANES 2005–2010",
    "10": "Genetic instruments for stool frequency",
    "11": "Mendelian randomization of genetically predicted stool frequency on eight outcomes",
    "12": f"Reverse Mendelian randomization: genetic liability to {PDn} and stool frequency",
    "13": f"Colocalisation of stool frequency and {PDn} at 12 loci",
    "14": "Bowel habits and quantitative susceptibility mapping and T2* of subcortical nuclei",
    "15": f"Stool form and bowel movement frequency fitted jointly for incident {PDn}",
    "16": "Bowel movement frequency and incident cardiometabolic disease, mental disorders and mortality",
    "17": f"Choice of time origin and the association of stool form with {PDn}",
    "18": f"Incremental predictive value of stool form for incident {PDn}",
    "19": f"Negative-control exposures from the Digestive Health Questionnaire and incident {PDn}",
    "20": f"Further sensitivity analyses for stool form and incident {PDn}",
    "21": f"Transit spectrum score and incident {PDn}",
    "22": f"Stool form and incident {PDn} by subgroup",
    "23": f"Absolute 5-year risk of {PDn} by stool form",
    "24": f"Stool form jointly with other prodromal markers and incident {PDn}",
    "25": f"Accelerometry sub-study: stool form, average acceleration and incident {PDn}",
}
D = MAP["data"]
NOTES = {
    "1": "",
    "2": ("Associations at FDR < 0.05 (model 2). Contrasts are versus one bowel movement per day (frequency), never (hard or "
          "loose stools) or normal bowel habits (phenotype)."),
    "3": f"As Supplementary Data {D['2']}, for digestive outcomes (ICD-10 chapter K).",
    "4a": "All 395 outcomes with at least 100 incident events and all ten contrasts (model 2).",
    "4b": f"As Supplementary Data {D['4a']}, adjusted for age and sex only (model 1).",
    "5": "Estimates for the model ladder, lag analyses, exclusions, follow-up windows and Fine–Gray competing-risk models.",
    "6": "",
    "7": "Model 2.",
    "8": "Survey-weighted models; mortality followed to 31 December 2019.",
    "9": "Survey-weighted logistic regression adjusted for age and sex among adults aged 40 years or older.",
    "10": "Genome-wide significant lead variants at least 10 Mb apart from the European-ancestry stool-frequency GWAS (GCST90565484).",
    "11": "Outcomes from FinnGen R12.",
    "12": f"Exposure: {PDn} GWAS of Nalls et al. (GCST009325) or FinnGen R12; outcome: stool frequency (GCST90565484).",
    "13": "Windows of ±500 kb around each stool-frequency lead variant; priors p1 = p2 = 1 × 10⁻⁴ and p12 = 1 × 10⁻⁵.",
    "14": "Standardized differences (SD units) from linear regression with robust standard errors.",
    "15": S15_NOTE,
    "16": "Events (rate per 1,000 person-years) by category; hazard ratios from model 2 with one bowel movement per day as reference.",
    "17": S17_NOTE, "18": S18_NOTE,
    "19": "Model 2. Reference: never (frequency) or no (current symptoms).",
    "20": ("Hazard ratios (95% CI), model 2 (multiple imputation and complete-case models include the same covariates with "
           "missing values imputed or excluded instead of median-filled). Reference: never."),
    "21": S21_NOTE,
    "22": "Model 2; P for interaction from likelihood-ratio tests of exposure-by-subgroup product terms.",
    "23": ("Crude: Aalen–Johansen cumulative incidence with death as a competing event. Standardized: mean predicted 5-year risk "
           "from model 2 with every participant assigned each category; 95% CIs from 200 bootstrap replicates."),
    "24": S24_NOTE, "25": S25_NOTE,
}
abbr_of = lambda old: "S" + old.rstrip("ab")
FIGDIR = P + "figures/npjPD_final/"
SUPPFIG = {1: "SuppFig1_subgroups.png", 2: "SuppFig2_frequency_other_outcomes.png", 3: "SuppFig3_genetics_imaging.png"}
legend_md = open(P + "manuscript/src_en/legends.md", encoding="utf-8").read().split("\n\n")
SUPPLEG = {int(re.match(r"\*\*Supplementary Fig\. (\d)", x).group(1)): x for x in legend_md if x.startswith("**Supplementary Fig.")}


def md_runs(par, text, size=9):
    for part in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|~[^~]+?~)", text):
        if not part:
            continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*"):
            r = par.add_run(part[1:-1]); r.italic = True
        elif part.startswith("~"):
            r = par.add_run(part[1:-1]); r.font.subscript = True
        else:
            r = par.add_run(part)
        r.font.size = Pt(size)


inv_t = {v: k for k, v in MAP["table"].items()}
inv_d = {v: k for k, v in MAP["data"].items()}
ms_title = open(P + "manuscript/src_en/title_abstract.md", encoding="utf-8").read().split("\n", 1)[0].lstrip("# ").strip()

# ---------------------------------------------------------------- Supplementary Information (Word)
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(9); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
sec0 = doc.sections[0]
for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(sec0, m, _Cm(2.0))
p = doc.add_paragraph(); r = p.add_run("Supplementary Information"); r.bold = True; r.font.size = Pt(16)
p = doc.add_paragraph(); r = p.add_run(ms_title); r.italic = True; r.font.size = Pt(11)
p = doc.add_paragraph(); r = p.add_run("Contents"); r.bold = True; r.font.size = Pt(11)
for n in sorted(SUPPLEG):
    t_ = re.match(r"\*\*Supplementary Fig\. \d (.+?)\*\*", SUPPLEG[n]).group(1).rstrip(".")
    doc.add_paragraph(f"Supplementary Fig. {n}  {t_}")
for n in sorted(inv_t):
    doc.add_paragraph(f"Supplementary Table {n}  {TITLES[inv_t[n]]}")
p = doc.add_paragraph(); r = p.add_run("Supplementary Data 1–10 (separate Excel file, Supplementary_Data.xlsx)"); r.bold = True
for n in sorted(inv_d):
    doc.add_paragraph(f"Supplementary Data {n}  {TITLES[inv_d[n]]}")
for n in sorted(SUPPLEG):
    doc.add_page_break()
    doc.add_picture(FIGDIR + SUPPFIG[n], width=_Cm(17))
    md_runs(doc.add_paragraph(), SUPPLEG[n])
for n in sorted(inv_t):
    old = inv_t[n]; d = S[SKEY[old]]
    add_table(f"Supplementary Table {n}. {TITLES[old]}", d.round(4) if hasattr(d, "round") else d, NOTES[old],
              landscape=d.shape[1] >= 7, abbr_key=abbr_of(old))
SI_OUT = P + "manuscript/Supplementary_Information.docx"
doc.save(SI_OUT)

# ---------------------------------------------------------------- Supplementary Data (Excel)
wb = Workbook(); ws0 = wb.active; ws0.title = "Contents"
ws0["A1"] = "Supplementary Data for: " + ms_title; ws0["A1"].font = Font(bold=True)
for i, n in enumerate(sorted(inv_d), start=3):
    ws0.cell(row=i, column=1, value=f"Supplementary Data {n}"); ws0.cell(row=i, column=2, value=TITLES[inv_d[n]])
ws0.column_dimensions["A"].width = 22; ws0.column_dimensions["B"].width = 90
for n in sorted(inv_d):
    old = inv_d[n]; d = S[SKEY[old]]
    ws = wb.create_sheet(f"Supplementary Data {n}")
    ws.cell(row=1, column=1, value=f"Supplementary Data {n}. {TITLES[old]}").font = Font(bold=True)
    for j, c in enumerate(d.columns, start=1):
        ws.cell(row=3, column=j, value=str(c)).font = Font(bold=True)
    for i, row in enumerate(d.itertuples(index=False), start=4):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, "item") else v))
    r0 = len(d) + 5
    if NOTES[old]:
        ws.cell(row=r0, column=1, value=NOTES[old]); r0 += 1
    ws.cell(row=r0, column=1, value="Abbreviations: " + ABBR[abbr_of(old)])
    for j, c in enumerate(d.columns, start=1):
        width = max([len(str(c))] + [len(str(v)) for v in d.iloc[:200, j - 1]]) + 2
        ws.column_dimensions[ws.cell(row=3, column=j).column_letter].width = min(60, max(10, width))
DATA_OUT = P + "manuscript/Supplementary_Data.xlsx"
wb.save(DATA_OUT)
print("submission package:", SI_OUT, "|", DATA_OUT)
