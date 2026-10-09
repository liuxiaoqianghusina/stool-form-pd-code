# Minimal Markdown -> Word converter for the manuscript draft (headings, paragraphs, lists, pipe tables,
# **bold**, *italic*, ~subscript~, and 【placeholder】 highlighted in yellow).
import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src, dst = sys.argv[1], sys.argv[2]
doc = Document()
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Cm(2.5); sec.top_margin = sec.bottom_margin = Cm(2.5)
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(6)
for lvl, size in [(1, 16), (2, 13), (3, 12)]:
    h = doc.styles[f"Heading {lvl}"]; h.font.name = "Times New Roman"; h.font.size = Pt(size); h.font.bold = True
    h.font.color.rgb = RGBColor(0, 0, 0); h.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

TOK = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|~[^~]+?~|\^[^^]+?\^|【[^】]*】)")
def add_runs(par, text):
    for part in TOK.split(text):
        if not part: continue
        if part.startswith("**"): r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*"): r = par.add_run(part[1:-1]); r.italic = True
        elif part.startswith("~"): r = par.add_run(part[1:-1]); r.font.subscript = True
        elif part.startswith("^"): r = par.add_run(part[1:-1]); r.font.superscript = True
        elif part.startswith("【"): r = par.add_run(part); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else: r = par.add_run(part.replace("\\*", "*"))
        r.font.name = "Times New Roman"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

lines = open(src, encoding="utf-8").read().split("\n")
i = 0
while i < len(lines):
    ln = lines[i]
    if not ln.strip(): i += 1; continue
    if ln.strip() == "\\pagebreak":
        doc.add_page_break(); i += 1; continue
    m = re.match(r"^(#{1,3})\s+(.*)", ln)
    if m:
        add_runs(doc.add_heading(level=len(m.group(1))), m.group(2)); i += 1; continue
    if ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip(" \t") for c in lines[i].strip().strip("|").split("|")]   # keep em-space indents
            if not all(re.fullmatch(r"-+", c) for c in cells): rows.append(cells)
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))   # three-line table: rules above/below header, below body
        for r, row in enumerate(rows):
            for c, txt in enumerate(row):
                cell = t.cell(r, c); cell.text = ""; p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
                add_runs(p, txt)
                for run in p.runs:
                    run.font.size = Pt(8.5); run.bold = run.bold or r == 0
                if c > 0:
                    p.alignment = 1
                edges = (("top", "bottom") if r == 0 else ()) + (("bottom",) if r == len(rows) - 1 else ())
                tcPr = cell._tc.get_or_add_tcPr(); bd = OxmlElement("w:tcBorders")
                for e in ("top", "left", "bottom", "right"):
                    el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single" if e in edges else "nil")
                    if e in edges:
                        el.set(qn("w:sz"), "8"); el.set(qn("w:color"), "000000")
                    bd.append(el)
                tcPr.append(bd)
        doc.add_paragraph(); continue
    m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", ln)
    if m:
        depth = len(m.group(1)) // 2
        style = "List Number" if m.group(2)[0].isdigit() else "List Bullet"
        if depth: style += f" {min(depth + 1, 3)}"
        add_runs(doc.add_paragraph(style=style), m.group(3)); i += 1; continue
    p = doc.add_paragraph(); add_runs(p, ln.strip()); i += 1
doc.save(dst)
print("saved", dst)
